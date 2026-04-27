#!/usr/bin/env python3
"""
Дамп SQLite бота: по умолчанию для отладки пишет Excel (удобно читать в LibreOffice/Excel).
В консоль — краткое подтверждение; таблицы в терминал: флаг --stdout.

Отключить Excel: вверху файла поставьте DUMP_TO_EXCEL = False.

Примеры:
  python dump_bot_db.py
  python dump_bot_db.py --limit 200
  python dump_bot_db.py --stdout
  python dump_bot_db.py --excel другой_путь.xlsx
  python dump_bot_db.py --no-excel --stdout
"""

from __future__ import annotations

import argparse
import os
import sqlite3
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
_PROMPT_BOT = Path(__file__).resolve().parent
_DEFAULT_DB = _PROMPT_BOT / "data" / "prompt_bot.sqlite"
_DEFAULT_XLSX = _PROMPT_BOT / "data" / "bot_dump.xlsx"

# --- отладка: Excel при каждом запуске. Потом поставьте False или закомментируйте вызов в main(). ---
DUMP_TO_EXCEL: bool = True
DUMP_EXCEL_PATH: Path | None = None  # None = путь по умолчанию data/bot_dump.xlsx


def _load_env() -> None:
    try:
        from dotenv import load_dotenv
    except ImportError:
        return
    load_dotenv(_ROOT / ".env")
    load_dotenv(_ROOT / ".env.local", override=True)
    load_dotenv(_PROMPT_BOT / ".env")
    load_dotenv(_PROMPT_BOT / ".env.local", override=True)


def _db_path() -> str:
    p = (os.getenv("VK_BOT_DB_PATH") or "").strip()
    if p:
        return p
    return str(_DEFAULT_DB)


def _short(s: str | None, n: int) -> str:
    if s is None:
        return ""
    s = s.replace("\n", " ").replace("\r", "")
    if len(s) <= n:
        return s
    return s[: n - 1] + "…"


def _cell_str(val: object, max_len: int = 32000) -> str:
    if val is None:
        return ""
    s = str(val)
    if len(s) > max_len:
        return s[: max_len - 12] + "…[truncated]"
    return s


def _print_table(
    conn: sqlite3.Connection,
    title: str,
    sql: str,
    params: tuple = (),
    *,
    wide_cols: tuple[int, ...] = (),
) -> int:
    cur = conn.execute(sql, params)
    rows = cur.fetchall()
    print(f"\n=== {title} ({len(rows)} строк) ===")
    if not rows:
        return 0
    cols = [d[0] for d in cur.description]
    widths = [max(len(c), *(len(_short(str(r[i]), 500)) for r in rows)) for i, c in enumerate(cols)]
    for i in wide_cols:
        if i < len(widths):
            widths[i] = min(widths[i], 72)
    header = " | ".join(c.ljust(widths[i]) for i, c in enumerate(cols))
    print(header)
    print("-+-".join("-" * w for w in widths))
    for r in rows:
        parts = []
        for i, c in enumerate(cols):
            val = r[i]
            text = "" if val is None else str(val)
            if i in wide_cols:
                text = _short(text, widths[i])
            parts.append(text.ljust(widths[i]))
        print(" | ".join(parts))
    return len(rows)


def _export_to_excel(
    conn: sqlite3.Connection,
    xlsx_path: Path,
    *,
    limit: int,
    table: str,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        print(
            "Нужен пакет openpyxl: pip install openpyxl",
            file=sys.stderr,
        )
        raise SystemExit(1) from exc

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    header_font = Font(bold=True)

    def add_sheet(name: str, sql: str, params: tuple) -> None:
        nonlocal first
        if first:
            ws = wb.active
            ws.title = name[:31]
            first = False
        else:
            ws = wb.create_sheet(title=name[:31])
        cur = conn.execute(sql, params)
        rows = cur.fetchall()
        cols = [d[0] for d in cur.description] if cur.description else []
        for c, h in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
        for ri, r in enumerate(rows, start=2):
            for ci, val in enumerate(r, start=1):
                ws.cell(row=ri, column=ci, value=_cell_str(val))
        for col in ws.columns:
            letter = col[0].column_letter
            maxlen = min(
                60,
                max(len(str(c.value or "")) for c in col[: min(200, len(col))]) + 2,
            )
            ws.column_dimensions[letter].width = maxlen

    lim = (limit,)
    if table in ("events", "all"):
        add_sheet(
            "bot_events",
            """
            SELECT id, created_at, vk_user_id, peer_id, kind, summary, payload_json
            FROM bot_events
            ORDER BY id DESC
            LIMIT ?
            """,
            lim,
        )
    if table in ("users", "all"):
        add_sheet(
            "users",
            """
            SELECT id, vk_user_id, peer_id, display_name, subscription_tier, subscription_status,
                   subscription_valid_until, prepaid_units_remaining, payment_provider,
                   payment_external_id, last_payment_at, last_payment_note, created_at, updated_at
            FROM users
            ORDER BY id DESC
            LIMIT ?
            """,
            lim,
        )
    if table in ("paid", "all"):
        add_sheet(
            "paid_request_units",
            """
            SELECT id, created_at, vk_user_id, peer_id, unit_type, settled, bot_event_id, metadata_json
            FROM paid_request_units
            ORDER BY id DESC
            LIMIT ?
            """,
            lim,
        )
    if table in ("ledger", "all"):
        add_sheet(
            "subscription_ledger",
            """
            SELECT id, created_at, vk_user_id, provider, external_id, amount_minor, currency,
                   status, plan_code, valid_until, raw_json
            FROM subscription_ledger
            ORDER BY id DESC
            LIMIT ?
            """,
            lim,
        )

    wb.save(xlsx_path)


def main() -> int:
    _load_env()
    parser = argparse.ArgumentParser(description="Дамп SQLite бота vk_prompt_bot")
    parser.add_argument(
        "--limit",
        type=int,
        default=50,
        metavar="N",
        help="сколько последних строк на лист (по умолчанию 50)",
    )
    parser.add_argument(
        "--table",
        choices=("events", "users", "paid", "ledger", "all"),
        default="all",
        help="какие таблицы выгрузить (по умолчанию all)",
    )
    parser.add_argument(
        "--path",
        type=str,
        default="",
        help="явный путь к .sqlite",
    )
    parser.add_argument(
        "--stdout",
        action="store_true",
        help="дополнительно печатать таблицы в терминал",
    )
    parser.add_argument(
        "--excel",
        type=str,
        nargs="?",
        const="__default__",
        default=None,
        metavar="FILE.xlsx",
        help="путь к Excel; без значения — data/bot_dump.xlsx",
    )
    parser.add_argument(
        "--no-excel",
        action="store_true",
        help="не писать Excel (только если включён DUMP_TO_EXCEL)",
    )
    args = parser.parse_args()

    path = args.path.strip() or _db_path()
    if not Path(path).is_file():
        print(f"Файл БД не найден: {path}", file=sys.stderr)
        return 1

    lim = max(1, min(args.limit, 5000))
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row

    do_excel = DUMP_TO_EXCEL and not args.no_excel
    excel_path: Path | None = None
    if do_excel:
        if args.excel is None or args.excel == "__default__":
            excel_path = DUMP_EXCEL_PATH if DUMP_EXCEL_PATH is not None else _DEFAULT_XLSX
        else:
            excel_path = Path(args.excel)

    try:
        print(f"БД: {path}")
        ver = conn.execute("SELECT value FROM schema_meta WHERE key = 'version'").fetchone()
        if ver:
            print(f"schema_meta.version = {ver[0]}")

        if do_excel and excel_path is not None:
            _export_to_excel(conn, excel_path, limit=lim, table=args.table)
            print(f"Excel записан: {excel_path.resolve()}")

        if args.stdout or not do_excel:
            if args.table in ("events", "all"):
                _print_table(
                    conn,
                    "bot_events (сначала новые)",
                    """
                    SELECT id, created_at, vk_user_id, peer_id, kind, summary, payload_json
                    FROM bot_events
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (lim,),
                    wide_cols=(5, 6),
                )
            if args.table in ("users", "all"):
                _print_table(
                    conn,
                    "users",
                    """
                    SELECT id, vk_user_id, peer_id, subscription_tier, subscription_status,
                           subscription_valid_until, prepaid_units_remaining,
                           created_at, updated_at
                    FROM users
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (lim,),
                )
            if args.table in ("paid", "all"):
                _print_table(
                    conn,
                    "paid_request_units",
                    """
                    SELECT id, created_at, vk_user_id, peer_id, unit_type, settled, bot_event_id, metadata_json
                    FROM paid_request_units
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (lim,),
                    wide_cols=(7,),
                )
            if args.table in ("ledger", "all"):
                _print_table(
                    conn,
                    "subscription_ledger",
                    """
                    SELECT id, created_at, vk_user_id, provider, external_id, status,
                           amount_minor, currency, plan_code, valid_until
                    FROM subscription_ledger
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (lim,),
                )
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
